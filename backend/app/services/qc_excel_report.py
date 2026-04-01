from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.config import BASE_DIR
from app.core.security import utc_now
from app.models.enums import QCCheckStatus, QCExecutionOutcome, TaskScope
from app.models.house import HouseType, PanelDefinition
from app.models.qc import (
    QCCheckDefinition,
    QCCheckInstance,
    QCExecution,
    QCFailureModeDefinition,
)
from app.models.stations import Station
from app.models.work import PanelUnit, WorkOrder, WorkUnit

try:  # pragma: no cover - import guard only
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    _OPENPYXL_ERROR: Exception | None = None
except Exception as exc:  # pragma: no cover - import guard only
    Workbook = None  # type: ignore[assignment]
    Image = None  # type: ignore[assignment]
    Alignment = None  # type: ignore[assignment]
    Border = None  # type: ignore[assignment]
    Font = None  # type: ignore[assignment]
    PatternFill = None  # type: ignore[assignment]
    Side = None  # type: ignore[assignment]
    _OPENPYXL_ERROR = exc


QC_REPORT_LOGO_PATH = BASE_DIR / "app" / "assets" / "reports" / "patagual-home-logo.png"
_SCOPE_LABELS: dict[TaskScope, str] = {
    TaskScope.MODULE: "Modulo",
    TaskScope.PANEL: "Panel",
    TaskScope.AUX: "Aux",
}
_SCOPE_SORT_ORDER: dict[TaskScope, int] = {
    TaskScope.MODULE: 0,
    TaskScope.PANEL: 1,
    TaskScope.AUX: 2,
}
_SHEET_TITLE_INVALID_CHARS = re.compile(r"[\[\]\*:/\\?]")
_NATURAL_TOKEN_RE = re.compile(r"(\d+)")


@dataclass(slots=True)
class PassedCheckReportRow:
    check_instance_id: int
    check_definition_id: int | None
    check_name: str
    module_number: int
    panel_code: str | None
    scope: TaskScope
    station_name: str | None
    performed_at: datetime
    notes: str | None


@dataclass(slots=True)
class HouseReportSheet:
    work_order_id: int
    project_name: str
    house_identifier: str | None
    house_type_name: str
    rows: list[PassedCheckReportRow] = field(default_factory=list)


def build_qc_dashboard_excel_report(db: Session) -> tuple[bytes, str]:
    if _OPENPYXL_ERROR is not None:
        raise RuntimeError("openpyxl is not available") from _OPENPYXL_ERROR

    generated_at = utc_now()
    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    houses = _load_house_report_sheets(db)
    failure_mode_lookup = _load_failure_mode_lookup(db, houses)

    if not houses:
        _build_empty_sheet(workbook, generated_at)
    else:
        used_titles: set[str] = set()
        for house in houses:
            _build_house_sheet(
                workbook,
                used_titles=used_titles,
                house=house,
                failure_mode_lookup=failure_mode_lookup,
                generated_at=generated_at,
            )

    output = BytesIO()
    workbook.save(output)
    filename = f"qc_dashboard_report_{generated_at.astimezone():%Y%m%d_%H%M%S}.xlsx"
    return output.getvalue(), filename


def _load_house_report_sheets(db: Session) -> list[HouseReportSheet]:
    latest_execution_subquery = (
        select(
            QCExecution.check_instance_id.label("check_instance_id"),
            QCExecution.outcome.label("outcome"),
            QCExecution.notes.label("notes"),
            QCExecution.performed_at.label("performed_at"),
            func.row_number()
            .over(
                partition_by=QCExecution.check_instance_id,
                order_by=(QCExecution.performed_at.desc(), QCExecution.id.desc()),
            )
            .label("row_number"),
        )
        .subquery()
    )

    rows = list(
        db.execute(
            select(
                WorkOrder.id.label("work_order_id"),
                WorkOrder.project_name.label("project_name"),
                WorkOrder.house_identifier.label("house_identifier"),
                HouseType.name.label("house_type_name"),
                WorkUnit.module_number.label("module_number"),
                QCCheckInstance.id.label("check_instance_id"),
                QCCheckInstance.check_definition_id.label("check_definition_id"),
                QCCheckInstance.scope.label("scope"),
                QCCheckInstance.ad_hoc_title.label("ad_hoc_title"),
                QCCheckDefinition.name.label("check_name"),
                Station.name.label("station_name"),
                PanelDefinition.panel_code.label("panel_code"),
                latest_execution_subquery.c.notes.label("execution_notes"),
                latest_execution_subquery.c.performed_at.label("performed_at"),
            )
            .select_from(QCCheckInstance)
            .join(WorkUnit, QCCheckInstance.work_unit_id == WorkUnit.id)
            .join(WorkOrder, WorkUnit.work_order_id == WorkOrder.id)
            .join(HouseType, WorkOrder.house_type_id == HouseType.id)
            .join(
                QCCheckDefinition,
                QCCheckInstance.check_definition_id == QCCheckDefinition.id,
                isouter=True,
            )
            .join(Station, QCCheckInstance.station_id == Station.id, isouter=True)
            .join(PanelUnit, QCCheckInstance.panel_unit_id == PanelUnit.id, isouter=True)
            .join(
                PanelDefinition,
                PanelUnit.panel_definition_id == PanelDefinition.id,
                isouter=True,
            )
            .join(
                latest_execution_subquery,
                latest_execution_subquery.c.check_instance_id == QCCheckInstance.id,
            )
            .where(QCCheckInstance.status == QCCheckStatus.CLOSED)
            .where(latest_execution_subquery.c.row_number == 1)
            .where(latest_execution_subquery.c.outcome == QCExecutionOutcome.PASS)
            .order_by(
                WorkOrder.project_name,
                WorkOrder.house_identifier,
                WorkUnit.module_number,
                PanelDefinition.panel_code,
                QCCheckInstance.id,
            )
        ).all()
    )

    houses_by_id: dict[int, HouseReportSheet] = {}
    for row in rows:
        house = houses_by_id.get(row.work_order_id)
        if house is None:
            house = HouseReportSheet(
                work_order_id=row.work_order_id,
                project_name=row.project_name or "",
                house_identifier=row.house_identifier,
                house_type_name=row.house_type_name or "",
            )
            houses_by_id[row.work_order_id] = house

        check_name = (row.check_name or row.ad_hoc_title or f"Check #{row.check_instance_id}").strip()
        house.rows.append(
            PassedCheckReportRow(
                check_instance_id=row.check_instance_id,
                check_definition_id=row.check_definition_id,
                check_name=check_name,
                module_number=row.module_number,
                panel_code=row.panel_code,
                scope=row.scope,
                station_name=row.station_name,
                performed_at=row.performed_at,
                notes=row.execution_notes,
            )
        )

    houses = list(houses_by_id.values())
    houses.sort(
        key=lambda house: (
            house.project_name.lower(),
            (house.house_identifier or "").lower(),
            house.work_order_id,
        )
    )
    for house in houses:
        house.rows.sort(
            key=lambda item: (
                item.module_number,
                _SCOPE_SORT_ORDER.get(item.scope, 99),
                _natural_sort_key(item.panel_code),
                item.check_name.lower(),
                item.performed_at,
                item.check_instance_id,
            )
        )
    return houses


def _load_failure_mode_lookup(
    db: Session,
    houses: list[HouseReportSheet],
) -> dict[int | None, list[str]]:
    check_definition_ids = sorted(
        {
            row.check_definition_id
            for house in houses
            for row in house.rows
            if row.check_definition_id is not None
        }
    )

    stmt = (
        select(QCFailureModeDefinition)
        .where(QCFailureModeDefinition.active.is_(True))
        .order_by(QCFailureModeDefinition.check_definition_id, QCFailureModeDefinition.name)
    )
    if check_definition_ids:
        stmt = stmt.where(
            (QCFailureModeDefinition.check_definition_id.in_(check_definition_ids))
            | (QCFailureModeDefinition.check_definition_id.is_(None))
        )
    else:
        stmt = stmt.where(QCFailureModeDefinition.check_definition_id.is_(None))

    mode_rows = list(db.execute(stmt).scalars())
    grouped: dict[int | None, list[str]] = defaultdict(list)
    for mode in mode_rows:
        grouped[mode.check_definition_id].append(mode.name)

    return {key: _dedupe_preserving_order(values) for key, values in grouped.items()}


def _build_empty_sheet(workbook: Workbook, generated_at: datetime) -> None:
    sheet = workbook.create_sheet(title="Resumen")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 72
    sheet.merge_cells("B2:C3")
    sheet["B2"] = "CHECK LIST QC - INSPECCIONES APROBADAS"
    sheet["B2"].font = Font(bold=True, size=14)
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["B5"] = "Exportado"
    sheet["C5"] = _format_datetime(generated_at)
    sheet["B7"] = "Estado"
    sheet["C7"] = "No hay inspecciones QC cerradas con resultado aprobado."


def _build_house_sheet(
    workbook: Workbook,
    *,
    used_titles: set[str],
    house: HouseReportSheet,
    failure_mode_lookup: dict[int | None, list[str]],
    generated_at: datetime,
) -> None:
    title_seed = house.house_identifier or f"Casa {house.work_order_id}"
    sheet_title = _make_unique_sheet_title(f"Casa {title_seed}", used_titles)
    ws = workbook.create_sheet(title=sheet_title)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B10"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    widths = {
        "B": 10,
        "C": 12,
        "D": 11,
        "E": 18,
        "F": 34,
        "G": 34,
        "H": 18,
        "I": 7,
        "J": 7,
        "K": 7,
        "L": 28,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.row_dimensions[2].height = 28
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[6].height = 22
    ws.row_dimensions[7].height = 22
    ws.row_dimensions[9].height = 22
    ws.row_dimensions[10].height = 24

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    header_fill = PatternFill(fill_type="solid", fgColor="4F81BD")
    section_fill = PatternFill(fill_type="solid", fgColor="D9E6F2")
    note_fill = PatternFill(fill_type="solid", fgColor="F8FAFC")
    pass_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    white_font = Font(bold=True, color="FFFFFF")
    header_font = Font(bold=True)
    title_font = Font(bold=True, size=14)

    if QC_REPORT_LOGO_PATH.exists():
        try:
            logo = Image(str(QC_REPORT_LOGO_PATH))
            logo.width = 118
            logo.height = 67
            ws.add_image(logo, "B2")
        except Exception:
            pass

    ws.merge_cells("E2:L3")
    ws["E2"] = "CHECK LIST QC - INSPECCIONES APROBADAS"
    ws["E2"].alignment = center
    ws["E2"].font = title_font
    ws["E2"].fill = header_fill
    ws["E2"].border = border

    ws["E4"] = "Reporte exportado desde dashboard QC"
    ws["E4"].alignment = left
    ws["E4"].font = Font(italic=True, size=10)

    unique_modules = sorted({row.module_number for row in house.rows})
    expanded_row_count = sum(
        max(1, len(_resolve_failure_modes(failure_mode_lookup, row.check_definition_id)))
        for row in house.rows
    )

    _write_key_value(ws, "B6", "C6", "CLIENTE", "Patagual Home", border, header_fill, white_font, center, left)
    _write_key_value(
        ws,
        "E6",
        "F6",
        "PROYECTO",
        house.project_name or "-",
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "H6",
        "I6",
        "N° CASA",
        house.house_identifier or f"#{house.work_order_id}",
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "K6",
        "L6",
        "MODELO",
        house.house_type_name or "-",
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "B7",
        "C7",
        "MODULOS",
        ", ".join(str(module) for module in unique_modules) or "-",
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "E7",
        "F7",
        "CHECKS",
        str(len(house.rows)),
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "H7",
        "I7",
        "MODOS",
        str(expanded_row_count),
        border,
        header_fill,
        white_font,
        center,
        left,
    )
    _write_key_value(
        ws,
        "K7",
        "L7",
        "EXPORTADO",
        _format_datetime(generated_at),
        border,
        header_fill,
        white_font,
        center,
        left,
    )

    ws.merge_cells("B9:L9")
    ws["B9"] = "DETALLE DE INSPECCIONES APROBADAS"
    ws["B9"].alignment = left
    ws["B9"].font = header_font
    ws["B9"].fill = section_fill
    ws["B9"].border = border

    headers = [
        ("B10", "Modulo"),
        ("C10", "Panel"),
        ("D10", "Alcance"),
        ("E10", "Estacion"),
        ("F10", "Check aprobado"),
        ("G10", "Modo de falla"),
        ("H10", "Fecha QC"),
        ("I10", "Si"),
        ("J10", "No"),
        ("K10", "N/A"),
        ("L10", "Observaciones"),
    ]
    for coordinate, label in headers:
        ws[coordinate] = label
        ws[coordinate].alignment = center
        ws[coordinate].font = white_font
        ws[coordinate].fill = header_fill
        ws[coordinate].border = border

    row_cursor = 11
    current_module: int | None = None
    for item in house.rows:
        failure_modes = _resolve_failure_modes(failure_mode_lookup, item.check_definition_id)
        if current_module != item.module_number:
            current_module = item.module_number
            ws.merge_cells(f"B{row_cursor}:L{row_cursor}")
            ws[f"B{row_cursor}"] = f"Modulo {item.module_number}"
            ws[f"B{row_cursor}"].alignment = left
            ws[f"B{row_cursor}"].font = header_font
            ws[f"B{row_cursor}"].fill = section_fill
            ws[f"B{row_cursor}"].border = border
            row_cursor += 1

        for index, failure_mode in enumerate(failure_modes):
            ws[f"B{row_cursor}"] = item.module_number if index == 0 else ""
            ws[f"C{row_cursor}"] = item.panel_code or "-"
            ws[f"D{row_cursor}"] = _SCOPE_LABELS.get(item.scope, item.scope.value)
            ws[f"E{row_cursor}"] = item.station_name or "Sin estacion"
            ws[f"F{row_cursor}"] = item.check_name
            ws[f"G{row_cursor}"] = failure_mode
            ws[f"H{row_cursor}"] = _format_datetime(item.performed_at)
            ws[f"I{row_cursor}"] = "x"
            ws[f"J{row_cursor}"] = ""
            ws[f"K{row_cursor}"] = ""
            ws[f"L{row_cursor}"] = item.notes or ""

            for column in "BCDEFGHIJKL":
                cell = ws[f"{column}{row_cursor}"]
                cell.border = border
                cell.alignment = center if column in {"B", "C", "D", "H", "I", "J", "K"} else left
                if column == "I":
                    cell.fill = pass_fill
                elif row_cursor % 2 == 0:
                    cell.fill = note_fill
            row_cursor += 1


def _write_key_value(
    ws,
    label_cell: str,
    value_cell: str,
    label: str,
    value: str,
    border,
    label_fill,
    label_font,
    label_alignment,
    value_alignment,
) -> None:
    ws[label_cell] = label
    ws[label_cell].border = border
    ws[label_cell].fill = label_fill
    ws[label_cell].font = label_font
    ws[label_cell].alignment = label_alignment
    ws[value_cell] = value
    ws[value_cell].border = border
    ws[value_cell].alignment = value_alignment


def _resolve_failure_modes(
    failure_mode_lookup: dict[int | None, list[str]],
    check_definition_id: int | None,
) -> list[str]:
    modes = list(failure_mode_lookup.get(check_definition_id, []))
    if check_definition_id is not None:
        modes.extend(failure_mode_lookup.get(None, []))
    else:
        modes = failure_mode_lookup.get(None, [])
    deduped = _dedupe_preserving_order(modes)
    return deduped or ["Sin modos de falla configurados"]


def _dedupe_preserving_order(values: list[str]) -> list[str]:
    deduped: list[str] = []
    seen: set[str] = set()
    for value in values:
        normalized = value.strip()
        if not normalized:
            continue
        key = normalized.casefold()
        if key in seen:
            continue
        seen.add(key)
        deduped.append(normalized)
    return deduped


def _make_unique_sheet_title(seed: str, used_titles: set[str]) -> str:
    base = _SHEET_TITLE_INVALID_CHARS.sub(" ", seed).strip()
    base = re.sub(r"\s+", " ", base) or "Casa"
    candidate = base[:31]
    if candidate not in used_titles:
        used_titles.add(candidate)
        return candidate

    suffix = 2
    while True:
        trailer = f" ({suffix})"
        trimmed = base[: max(1, 31 - len(trailer))].rstrip()
        candidate = f"{trimmed}{trailer}"
        if candidate not in used_titles:
            used_titles.add(candidate)
            return candidate
        suffix += 1


def _natural_sort_key(value: str | None) -> tuple[object, ...]:
    raw = (value or "").strip()
    if not raw:
        return ("",)
    parts = _NATURAL_TOKEN_RE.split(raw)
    tokens: list[object] = []
    for part in parts:
        if not part:
            continue
        if part.isdigit():
            tokens.append(int(part))
        else:
            tokens.append(part.lower())
    return tuple(tokens)


def _format_datetime(value: datetime) -> str:
    return value.astimezone().strftime("%d-%m-%Y %H:%M")
